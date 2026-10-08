"""Admin-only orders → .xlsx export (GET /api/admin/orders/export).

Kept separate from routers/admin.py for the same reason catalog_pdf.py is its
own module — document generation is a distinct concern from request routing.
"""
from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

STATUS_LABELS = {
    "pending": "Новый",
    "processing": "В обработке",
    "completed": "Выполнен",
    "cancelled": "Отменён",
}

HEADERS = [
    "Номер заказа", "Дата", "Клиент", "Телефон", "Email", "Адрес",
    "Статус", "Товары", "Сумма, сум", "Сумма, $",
]


def _items_summary(order) -> str:
    return "; ".join(
        f"{i.product_name_snapshot} ×{i.quantity}" for i in order.items
    )


def build_orders_xlsx(orders: list, usd_rate: Decimal | None) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Заказы"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F2937")  # gray-800, matches admin UI
    ws.append(HEADERS)
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center")

    for order in orders:
        total_uzs = Decimal(order.total_uzs)
        total_usd = (total_uzs / usd_rate) if usd_rate else None
        ws.append([
            order.order_number,
            order.created_at.strftime("%d.%m.%Y %H:%M") if order.created_at else "",
            order.customer_name,
            order.customer_phone,
            order.customer_email or "",
            order.customer_address or "",
            STATUS_LABELS.get(order.status, order.status),
            _items_summary(order),
            float(total_uzs),
            float(total_usd) if total_usd is not None else None,
        ])

    # Column widths — generous enough for real data, not autosizing (openpyxl
    # can't measure rendered text width without a spreadsheet engine).
    widths = [16, 16, 22, 16, 22, 28, 14, 50, 14, 12]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    # Number formatting for the two money columns
    money_cols = (9, 10)
    for row in ws.iter_rows(min_row=2, min_col=1, max_col=len(HEADERS)):
        for col_idx in money_cols:
            cell = row[col_idx - 1]
            if cell.value is not None:
                cell.number_format = "#,##0.00"

    ws.freeze_panes = "A2"

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
